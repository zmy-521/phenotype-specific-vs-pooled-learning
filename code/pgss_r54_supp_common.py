# -*- coding: utf-8 -*-
r"""R54 补充材料汇编的公共部分：来源读取、英文对照、数字格式、三线表写入 Word。只读冻结表，不计算新统计量。"""
import re, json, pathlib, hashlib
import numpy as np, pandas as pd

HERE = pathlib.Path(__file__).resolve().parent
ROOT = HERE.parent
SRC = {
    "R47": ROOT / "29_R47_写作交接" / "R47_冻结数字包.xlsx",
    "R48": ROOT / "30_R48_图稿与补充材料" / "R48_图源与补充材料.xlsx",
    "R49": ROOT / "31_R49_只读补齐" / "R49_只读补齐_证据表.xlsx",
    "R51": ROOT / "34_R51_基线只读提数" / "R51_Table1_Baseline_Source.xlsx",
    "R52": ROOT / "35_R52_NHANES_SCr2005_06_SENS" / "R52_NHANES_SCr_SENS_EVIDENCE.xlsx",
}
R50DIR = ROOT / "32_R50_讨论只读计数"
NHCOH = ROOT / "25_R44_NHANES跨领域复现" / "data" / "COHORT.json"
V5 = ROOT / "19_R38" / "data_v5"
_cache = {}


def sheet(src, name):
    k = (src, name)
    if k not in _cache:
        _cache[k] = pd.read_excel(SRC[src], sheet_name=name)
    return _cache[k].copy()


def sha16(p):
    return hashlib.sha256(pathlib.Path(p).read_bytes()).hexdigest()[:16]


# ── 英文对照 ────────────────────────────────────────────────────
DB = {"mimic": "MIMIC-IV", "eicu": "eICU-CRD", "nwicu": "NWICU", "nhanes": "NHANES", "MIMIC-IV": "MIMIC-IV", "eICU-CRD": "eICU-CRD",
      "NWICU": "NWICU", "NHANES": "NHANES", "nhanes（C1 pop，追加对照）": "NHANES (C1 sample-weighted, objective-aligned)"}
OC = {"OUT_HOSP_MORT": "In-hospital death", "OUT_LOS72": "ICU stay >72 h", "O_MORT10": "10-y all-cause death", "O_CVD10": "10-y CVD death"}
PH = {"女性": "Female", "年龄≥65": "Age ≥65 y", "贫血": "Anemia", "eGFR降低": "Reduced eGFR", "低白蛋白": "Low albumin", "肥胖": "Obesity"}
SC = {"A": "Track A", "B": "Track B", "bal": "Balanced", "pop": "Sample-weighted", "两轨等权": "Balanced", "全体": "Whole task"}
TS = {"排除实质收益": "Excluded", "证据不足": "Inconclusive", "支持实质收益": "Supported",
      "排除（收益<m）": "Excluded", "支持（收益>m）": "Supported"}
D0 = {"全负": "Entirely <0", "跨零": "Crosses 0", "全正": "Entirely >0",
      "区间整体 < 0": "Entirely <0", "区间整体 > 0": "Entirely >0", "区间跨 0": "Crosses 0", "全部重复为 0": "All replicates 0",
      "区间整体 < m": "Entirely <m", "区间整体 ≥ m": "Entirely ≥m", "区间跨 m": "Crosses m"}
FEAT = {"MCV": "MCV", "RDW": "RDW", "尿素氮": "Blood urea nitrogen", "平均动脉压": "Mean arterial pressure", "年龄": "Age", "心率": "Heart rate",
        "收缩压": "Systolic blood pressure", "白细胞计数": "White blood cell count", "碳酸氢盐": "Bicarbonate", "舒张压": "Diastolic blood pressure",
        "血小板": "Platelet count", "血氯": "Chloride", "血糖": "Glucose", "血红蛋白": "Hemoglobin", "血细胞比容": "Hematocrit", "血肌酐": "Serum creatinine",
        "血钙": "Calcium", "血钠": "Sodium", "血钾": "Potassium", "阴离子间隙": "Anion gap", "性别F": "Female sex", "SpO2": "SpO2", "呼吸频率": "Respiratory rate",
        "白蛋白": "Albumin", "身高cm": "Height", "体重kg": "Weight", "BMI": "Body mass index", "体温": "Temperature", "eGFR_landmark": "eGFR at landmark"}
MISC = {"main": "Main", "cert": "Certification", "lc": "Learning curve", "HTFS面板": "HTFS panel", "P*选择": "P* selection", "P*_q": "P*_objective",
        "全体": "All", "认证共同对象": "Certification common set", "True": "Yes", "False": "No", "是": "Yes", "否": "No",
        "NHANES O_CVD10（删除 1 人）": "NHANES 10-y CVD death (1 participant removed)",
        "NWICU 住院死亡（删除 200 人次，约 1%）": "NWICU in-hospital death (200 stays removed, about 1%)",
        "R43 B1/B2（评价对象 = 全体）": "R43 B1/B2 (all evaluated units)", "NWICU v2（R46）": "NWICU v2 (R46)", "NHANES v2（R46）": "NHANES v2 (R46)",
        "O-2：{P*_D, HTFS_D}": "O-2: {P*_D, HTFS_D}", "O-7：全部 D 臂": "O-7: all D arms"}
COL = {  # 列名
    "库": "Database", "数据库": "Database", "结局": "Outcome", "表型": "Phenotype", "指标": "Metric", "对照": "Contrast", "参照": "Reference", "候选": "Candidate",
    "尺度": "Scale", "观察绝对差": "Absolute difference", "观察相对改善": "Relative improvement", "R参照": "Reference risk", "相对_P2_5": "Relative 2.5th pct",
    "相对_P97_5": "Relative 97.5th pct", "绝对_P2_5": "Absolute 2.5th pct", "绝对_P97_5": "Absolute 97.5th pct", "缺失replicate": "Missing replicates",
    "三态_m0.005": "Three-state (m=0.5%)", "三态_m0.010": "Three-state (m=1%)", "三态_m0.020": "Three-state (m=2%)", "三态_m1": "Three-state (m=1%)",
    "相对0方向": "Interval vs 0", "任务": "Tasks", "相对改善中位": "Median relative improvement", "相对改善最小": "Minimum relative improvement",
    "相对改善最大": "Maximum relative improvement", "绝对差中位": "Median absolute difference", "相对上界最大": "Maximum relative upper bound",
    "全负": "Entirely <0", "跨零": "Crosses 0", "全正": "Entirely >0", "排除m1": "Excluded (m=1%)", "证据不足m1": "Inconclusive (m=1%)", "支持m1": "Supported (m=1%)",
    "排除": "Excluded", "证据不足": "Inconclusive", "支持": "Supported", "点估计为负": "Negative point estimates",
    "轨": "Track", "θ观察": "θ observed", "θ重抽SD": "θ resampling SD", "max-T临界值": "max-T critical value", "max_T临界值": "max-T critical value",
    "同时下界": "Simultaneous lower", "同时上界": "Simultaneous upper", "零方差单列": "Zero-variance column", "同时判定_相对m": "Simultaneous judgment (m=1%)",
    "项数": "Items", "证据不足项": "Inconclusive items", "分量数": "Components", "绝对差": "Absolute difference", "t下界": "t lower", "t上界": "t upper",
    "相对差": "Relative difference", "相对t下界": "Relative t lower", "相对t上界": "Relative t upper", "分支": "Branch", "臂": "Arm", "选择": "Selection",
    "单元数": "Units", "Pstar": "Selection", "项": "Item", "目标": "Objective", "seed": "Seed", "范围": "Scope", "评价对象": "Evaluation set", "n": "n",
    "事件": "Events", "Brier": "Brier", "logloss": "Log loss", "平均预测": "Mean prediction", "观察事件率": "Observed event rate", "AUROC": "AUROC", "AUPRC": "AUPRC",
    "校准截距": "Calibration intercept", "校准斜率": "Calibration slope", "seed数": "Seeds", "平均预测减观察": "Mean prediction − observed",
    "任务乘seed": "Task×seed units", "main_units": "Main units", "cert_units": "Certification units", "认证比较纳入单位比例": "Proportion of units certified",
    "单元": "Units", "专业化": "Specialist enabled", "测试折实现达m": "Test-fold gain ≥m", "单元×轨": "Unit×track decisions", "外层折": "Outer fold",
    "认证参照": "Certification reference", "认证p": "Certification p", "认证SE": "Certification SE", "认证簇数": "Certification clusters",
    "测试折实现相对收益": "Test-fold realized relative gain", "评价人数": "Evaluated n", "纳入比例": "Proportion included", "R_PstarD": "R(P*_D)",
    "R_PstarF": "R(P*_F)", "R_CERT": "R(CERT)", "样本成本项": "Sample-cost term", "认证启用项": "Enabling term", "合计": "Total", "恒等式残差": "Identity residual",
    "样本成本项_相对R_PstarD": "Sample-cost term / R(P*_D)", "认证启用项_相对R_PstarD": "Enabling term / R(P*_D)", "合计_相对R_PstarD": "Total / R(P*_D)",
    "样本成本项相对中位": "Median sample-cost term (relative)", "认证启用项相对中位": "Median enabling term (relative)", "合计相对中位": "Median total (relative)",
    "认证启用项绝对最大": "Maximum |enabling term|", "恒等式残差绝对最大": "Maximum |identity residual|", "比例": "Target training share",
    "目标比例": "Target training share", "完整度": "Completeness", "已运行": "Units run", "实际比例最小": "Actual share minimum",
    "实际比例中位": "Actual share median", "实际比例最大": "Actual share maximum", "子集簇数中位": "Median subset clusters",
    "已运行实际比例中位": "Median actual share (units run)", "任务数": "Tasks", "全部外层折可评价任务": "Tasks evaluable in all outer folds",
    "HTFS_vs_P星_bal中位_全可评价任务": "Median balanced HTFS vs P* (fully evaluable tasks)", "有任一折可评价任务": "Tasks evaluable in any fold",
    "同组任务": "Tasks in group", "P星两轨等权中位": "Median balanced risk, P*", "HTFS两轨等权中位": "Median balanced risk, HTFS",
    "Hfull两轨等权中位": "Median balanced risk, Hfull", "P星下降": "Decrease, P*", "HTFS下降": "Decrease, HTFS", "HTFS下降更多": "HTFS decreased more",
    "条件": "Condition", "选择器": "Selector", "选择器甲": "Selector 1", "选择器乙": "Selector 2", "重复": "Replicates", "独立重复": "Independent replicates",
    "T2配对差均值": "T2 paired difference, mean", "T2配对差SE": "T2 paired difference, SE", "T2下": "T2 lower", "T2上": "T2 upper", "T1均值差": "T1 mean difference",
    "T2推翻T1排序": "T2 reverses T1 ordering", "来源": "Source", "修正前": "Before correction", "修正后": "After correction", "变化": "Change",
    "绝对变化": "Absolute change", "原区间半宽": "Original interval half-width", "绝对变化_除以原半宽": "|Change| / half-width", "符号翻转": "Sign flips",
    "三态改变": "Three-state changes", "行": "Rows", "比值中位": "Median ratio", "比值最大": "Maximum ratio", "绝对变化中位": "Median absolute change",
    "绝对变化最大": "Maximum absolute change", "候选对": "Candidate pair", "真值相对收益均值": "True relative gain, mean", "SE": "SE", "下": "Lower (mean−1.96 SE)",
    "上": "Upper (mean+1.96 SE)", "均值±1.96SE方向": "Mean ±1.96 SE vs 0", "相对m方向": "Mean ±1.96 SE vs m", "P5": "P5", "中位": "Median", "P95": "P95",
    "两轨等权部署改善均值": "Balanced deployment improvement, mean", "点估计为负的重复": "Replicates with negative estimate",
    "部署相对PstarD改善均值": "Deployment improvement vs P*_D, mean", "部署实际变差_分子": "Actual deterioration, n",
    "部署实际变差_分母": "Actual deterioration, denominator", "部署实际变差_CP下": "Actual deterioration, CP lower", "部署实际变差_CP上": "Actual deterioration, CP upper",
    "部署变差边界不确定": "Deterioration near boundary", "P*_F相对P*_D改善均值": "P*_F vs P*_D improvement, mean", "最小": "Minimum", "最大": "Maximum",
    "≥1轨不达标启用_分子": "≥1 track enabled below m, n", "≥1轨不达标启用_分母": "≥1 track enabled below m, denominator",
    "≥1轨不达标启用_CP下": "≥1 track enabled below m, CP lower", "≥1轨不达标启用_CP上": "≥1 track enabled below m, CP upper",
    "≥1轨实际变差启用_分子": "≥1 track enabled with actual deterioration, n", "≥1轨实际变差启用_CP上": "≥1 track enabled with actual deterioration, CP upper",
    "开发集n": "Development n", "π_A": "π_A", "完成重复": "Completed replicates", "开发事件率均值": "Mean development event rate",
    "评价A轨平均p": "Mean true p, track A", "评价B轨平均p": "Mean true p, track B", "A轨Bayes风险": "Bayes risk, track A", "B轨Bayes风险": "Bayes risk, track B",
    "δ": "δ", "均值": "Mean", "两轨等权改善均值": "Balanced improvement, mean", "A轨改善均值": "Track A improvement, mean", "A轨SE": "Track A SE",
    "B轨改善均值": "Track B improvement, mean", "B轨SE": "Track B SE", "A轨取HTFS_D的重复": "Track A chose HTFS_D, replicates",
    "B轨取HTFS_D的重复": "Track B chose HTFS_D, replicates", "候选集": "Candidate set", "n恒定": "n constant across seeds",
    "R_Pstar": "R(P*)", "R_Pstar_q": "R(P*_objective)", "R_HTFS_q": "R(HTFS_objective)", "R_HTFS": "R(HTFS)", "v10差距": "Original gap",
    "收回": "Gap recovered", "收回比例_相对v10差距": "Proportion of original gap recovered", "判读尺度": "Judgment scale",
    "任务人数": "Task n", "A轨人数": "Track A n", "纳入": "Evaluable", "缺失单元": "Missing units", "A轨事件": "Track A events", "A轨占比": "Track A share",
    "A轨事件率": "Track A event rate", "B轨事件率": "Track B event rate", "切点": "Cut-off", "单位": "Unit", "eGFR标化体表面积_m2": "eGFR BSA (m²)",
    "变量": "Variable", "有值比例": "Observed proportion", "P1": "P1", "中位数": "Median", "P99": "P99", "步骤": "Step", "人次": "ICU stays", "剩余": "Remaining",
    "N": "N", "itemid": "itemid", "窗内记录": "Records in window", "范围外剔除": "Out-of-range removed", "其中哨兵9999999": "Of which sentinel 9999999",
    "未配对分量": "Unpaired components", "仍并列且取值不同": "Remaining ties with different values", "有值就诊": "Stays with a value",
    "同charttime多来源聚合": "Multi-source aggregation at same charttime", "同charttime同itemid多版本": "Multiple versions, same charttime and itemid",
    "键相同但取值不同": "Same key, different values", "分组方式": "Grouping", "组": "Group", "窗内有体征": "Vital signs in window",
    "整段无任何体征记录": "No vital-sign records", "住院死亡率_窗内有体征": "In-hospital mortality, vital signs in window",
    "住院死亡率_窗内无体征": "In-hospital mortality, no vital signs in window", "R45任务人次": "R45 task stays", "排除人次": "Excluded stays",
    "排除_A轨": "Excluded, track A", "排除_B轨": "Excluded, track B", "v2任务人次": "v2 task stays", "R45事件": "R45 events", "v2事件": "v2 events",
    "v2_A轨事件": "v2 track A events", "v2_A轨人次": "v2 track A stays", "R44已纳入": "Included in R44", "未知者在任务人群中": "Unknown-cause death in task",
    "未知者所在轨": "Track of unknown-cause death", "需要重跑": "Rerun required", "包": "Package", "版本": "Version", "直接导入的文件数": "Files importing it",
    "作用": "Role", "B轨n": "Track B n", "B轨事件": "Track B events", "A轨n": "Track A n", "事件率": "Event rate", "分组单位": "Grouping unit",
}
for _m in ("AUROC", "AUPRC", "Brier", "logloss", "校准截距", "校准斜率"):
    _e = COL[_m]
    COL[f"{_m}_候选减参照"] = f"{_e} (candidate − reference)"
    COL[f"{_m}差中位"] = f"Median {_e} difference"
    COL[f"{_m}差为正"] = f"Positive {_e} differences"
    for _s, _t in (("_均值", "mean"), ("_最小", "minimum"), ("_最大", "maximum")):
        COL[f"{_m}{_s}"] = f"{_e}, {_t}"
        COL[f"{_m}_候选减参照{_s}"] = f"{_e} (candidate − reference), {_t}"
    for _s in ("count", "mean", "std", "min", "5%", "50%", "95%", "max"):
        COL[f"{_m}_{_s}"] = f"{_e}, {_s}"
for _m in ("Brier", "logloss"):
    COL[f"{_m}_相对改善"] = f"{COL[_m]} relative improvement"
    for _s, _t in (("_均值", "mean"), ("_最小", "minimum"), ("_最大", "maximum")):
        COL[f"{_m}_相对改善{_s}"] = f"{COL[_m]} relative improvement, {_t}"
for _a in ("A", "B", "bal"):
    _t = {"A": "track A", "B": "track B", "bal": "balanced"}[_a]
    COL[f"HTFS_vs_P*_{_a}_相对"] = f"HTFS vs P*, {_t} (relative)"
    COL[f"Hfull_vs_P*_{_a}_相对"] = f"Hfull vs P*, {_t} (relative)"
for _arm in ("P*", "HTFS", "Hfull"):
    for _a, _t in (("A", "track A"), ("B", "track B"), ("两轨等权", "balanced")):
        COL[f"{_arm}_{_a}"] = f"{_arm} risk, {_t}"
COL.update({"平均预测_均值": "Mean prediction, mean", "平均预测_最小": "Mean prediction, minimum", "平均预测_最大": "Mean prediction, maximum"})
COL.update({
    "任务患者数": "Distinct patients", "不可判定比例": "Not ascertainable proportion", "分母": "Denominator", "观测n": "Observed n", "缺失n": "Missing n",
    "缺失比例": "Missing proportion", "阳性n": "Positive n", "阳性比例": "Positive proportion", "上界最大": "Maximum upper bound",
    "A轨开发人数均值": "Mean track A development n", "A轨开发事件均值": "Mean track A development events", "B轨开发事件均值": "Mean track B development events",
    "评价10万": "Replicates finalized at 100k", "评价50万": "Replicates finalized at 500k", "评价100万": "Replicates finalized at 1M",
    "最终仍有近边界项的重复": "Replicates still near a boundary", "每重复秒均值": "Mean seconds per replicate", "拟合次数均值": "Mean fits per replicate",
    "峰值内存MB最大": "Peak memory MB (maximum)", "失败": "Failures",
    "两轨等权均值": "Balanced, mean", "两轨等权SE": "Balanced, SE", "两轨等权方向": "Balanced, mean ±1.96 SE vs 0", "两轨等权P5": "Balanced, P5", "两轨等权P95": "Balanced, P95",
    "两轨等权点估计为负": "Balanced, replicates < 0", "风险变化均值": "Mean risk change", "MIX更好重复": "Replicates with MIX better",
    "启用_分子": "Enabled, n", "启用_分母": "Enabled, denominator", "不达标启用_条件于确定不达标_分子": "Enabled given gain < m, n",
    "不达标启用_条件于确定不达标_分母": "Enabled given gain < m, denominator", "不达标启用_条件于确定不达标_CP上": "Enabled given gain < m, CP upper",
    "检出_条件于确定达标_分子": "Detected given gain ≥ m, n", "检出_条件于确定达标_分母": "Detected given gain ≥ m, denominator",
    "检出_条件于确定达标_CP下": "Detected given gain ≥ m, CP lower", "检出_条件于确定达标_CP上": "Detected given gain ≥ m, CP upper", "启用_m边界不确定": "Enabled near m boundary",
    "对象": "Object", "T1重复": "T1 replicates", "T1均值": "T1 mean", "T2重复": "T2 replicates", "T2均值": "T2 mean", "T2方向": "T2 mean ±1.96 SE vs 0",
    "窗内心率非缺失_比例": "Heart rate in window, proportion", "住院死亡率_窗内心率非缺失": "In-hospital mortality, heart rate in window",
    "住院死亡率_非窗内心率非缺失": "In-hospital mortality, heart rate not in window", "窗内任一规定体征非缺失_比例": "Any specified vital sign in window, proportion",
    "住院死亡率_窗内任一规定体征非缺失": "In-hospital mortality, any vital sign in window", "住院死亡率_非窗内任一规定体征非缺失": "In-hospital mortality, no vital sign in window"})
for _t, _e in (("A", "Track A"), ("B", "Track B")):
    COL.update({f"{_t}轨均值": f"{_e}, mean", f"{_t}轨方向": f"{_e}, mean ±1.96 SE vs 0", f"{_t}轨P5": f"{_e}, P5", f"{_t}轨P95": f"{_e}, P95",
                f"{_t}轨实际变差": f"{_e}, actual deterioration", f"{_t}轨重复": f"{_e}, replicates", f"{_t}轨实际变差_CP上": f"{_e}, actual deterioration CP upper"})
for _k, _e in (("X绝对值最大_中位", "max |X|, median"), ("X绝对值最大_最大", "max |X|, maximum"), ("X绝对值P99_9_均值", "P99.9 of |X|, mean"),
               ("X绝对值超4比例_均值", "proportion |X| > 4, mean"), ("线性预测子P0_1_均值", "linear predictor P0.1, mean"), ("线性预测子P99_9_均值", "linear predictor P99.9, mean"),
               ("线性预测子最小", "linear predictor minimum"), ("线性预测子最大", "linear predictor maximum"), ("p小于1e6比例_均值", "proportion p < 1e-6, mean"),
               ("p大于1减1e6比例_均值", "proportion p > 1 − 1e-6, mean")):
    COL[f"初始10万评价样本_{_k}"] = f"First 100k evaluation draws: {_e}"
MISC.update({"患者": "Patient", "任务内患者—医院连通分量": "Patient–hospital connected component within task", "个体（SEQN）": "Individual (SEQN)",
             "平均动脉压（计算值）": "Mean arterial pressure (calculated)", "chartevents 同设备 S/D 配对": "chartevents, same-device systolic/diastolic pairing",
             "chartevents 基线体型": "chartevents, baseline body size", "全部": "All", "k=23（规格即 S0）": "k=23 (specification equals S0)",
             "k=23（规格即 SZ）": "k=23 (specification equals SZ)", "并集=全部 23 列": "Union = all 23 columns", "HTFS_q 版本": "HTFS_objective version",
             "P*（v1.0，同 11 任务）": "P* (original, same 11 tasks)", "MIMIC-IV + eICU-CRD（= R43 B3）": "MIMIC-IV + eICU-CRD", "MIMIC-IV + eICU-CRD（= R43 B4）": "MIMIC-IV + eICU-CRD",
             "NWICU（v2）": "NWICU", "NHANES（v2）": "NHANES", "nhanes（目标一致性追加臂）": "NHANES (objective-aligned arms)",
             "MIX − POOLMIX（两轨等权，同重复配对）": "MIX − POOLMIX (balanced, paired within replicate)", "候选对真值相对收益": "Candidate-pair true relative gain",
             "选择器两轨等权部署改善": "Selector balanced deployment improvement"})

CJK = re.compile(r"[㐀-鿿（）；，：、]")


def tr(v):
    """单元格取值的英文化：精确对照优先，再做表型/分隔符等成分替换。数字原样返回。"""
    if v is None or (isinstance(v, float) and np.isnan(v)):
        return v
    if isinstance(v, (bool, np.bool_)):
        return "Yes" if v else "No"
    if isinstance(v, (int, float, np.integer, np.floating)):
        return v
    s = str(v)
    for m in (DB, OC, PH, SC, TS, D0, MISC, FEAT):
        if s in m:
            return m[s]
    t = s
    for k, e in OC.items():
        t = t.replace(k, e)
    for k, e in {**PH, **FEAT}.items():
        t = t.replace(k, e)
    t = t.replace("；", "; ").replace("，", ", ").replace("：", ": ").replace("（", " (").replace("）", ")").replace("×", " × ")
    return re.sub(r"\s+", " ", t).strip()


def english(df, drop=()):
    """长表英文化：去掉指定的中文自由文本列，列名与取值全部映射。"""
    d = df.drop(columns=[c for c in drop if c in df.columns]).copy()
    for c in d.columns:
        if d[c].dtype == object or d[c].dtype == bool:
            d[c] = d[c].map(tr)
    d.columns = [COL.get(c, tr(c)) for c in d.columns]
    return d


def residual_cjk(df):
    bad = [c for c in df.columns if CJK.search(str(c))]
    cells = int(sum(df[c].astype(str).str.contains(CJK).sum() for c in df.columns if df[c].dtype == object))
    return bad, cells


# ── 数字格式 ────────────────────────────────────────────────────
def fnum(v, dp=0):
    if v is None or (isinstance(v, float) and np.isnan(v)):
        return "—"
    v = float(v)
    s = f"{int(round(v)):,}" if dp == 0 else f"{v:,.{dp}f}"
    return s.replace("-", "−")


def fpct(v, dp=2, sign=False):
    """比例 → 百分数（不带 % 号，列名注明 %）。"""
    if v is None or (isinstance(v, float) and np.isnan(v)):
        return "—"
    s = f"{100 * float(v):+.{dp}f}" if sign else f"{100 * float(v):.{dp}f}"
    return s.replace("-", "−")


def fci(lo, hi, dp=2):
    return f"{fpct(lo, dp, True)} to {fpct(hi, dp, True)}"


def fnd(n, d):
    return f"{fnum(n)}/{fnum(d)}"


# ── Word 三线表 ─────────────────────────────────────────────────
def three_line(table):
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge, sz in (("top", "12"), ("bottom", "12"), ("left", None), ("right", None), ("insideH", None), ("insideV", None)):
        el = OxmlElement(f"w:{edge}")
        if sz:
            el.set(qn("w:val"), "single"); el.set(qn("w:sz"), sz); el.set(qn("w:space"), "0"); el.set(qn("w:color"), "000000")
        else:
            el.set(qn("w:val"), "nil")
        borders.append(el)
    tblPr.append(borders)
    for cell in table.rows[0].cells:  # 表头下线
        tcPr = cell._tc.get_or_add_tcPr()
        b = OxmlElement("w:tcBorders"); bt = OxmlElement("w:bottom")
        bt.set(qn("w:val"), "single"); bt.set(qn("w:sz"), "6"); bt.set(qn("w:space"), "0"); bt.set(qn("w:color"), "000000")
        b.append(bt); tcPr.append(b)


def add_table(doc, title, df, legend, font_pt=7.5):
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    p = doc.add_paragraph(); r = p.add_run(title); r.bold = True; r.font.size = Pt(10)
    t = doc.add_table(rows=1, cols=len(df.columns))
    for j, c in enumerate(df.columns):
        cell = t.rows[0].cells[j]; cell.text = ""
        run = cell.paragraphs[0].add_run(str(c)); run.bold = True; run.font.size = Pt(font_pt)
    for _, row in df.iterrows():
        cells = t.add_row().cells
        for j, v in enumerate(row):
            cells[j].text = ""
            run = cells[j].paragraphs[0].add_run("" if v is None or (isinstance(v, float) and np.isnan(v)) else str(v))
            run.font.size = Pt(font_pt)
            if j > 0:
                cells[j].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    three_line(t)
    # 列宽按内容长度分配（横向 A4 可用约 26 cm），不再平均分配
    from docx.shared import Cm
    w = []
    for j, c in enumerate(df.columns):
        lens = [len(str(c)) * 0.7] + [len(str(v)) for v in df.iloc[:, j].tolist()]
        w.append(min(max(max(lens), 4), 60))
    tot = sum(w)
    for j, wj in enumerate(w):
        for row in t.rows:
            row.cells[j].width = Cm(26.0 * wj / tot)
    for line in legend:
        q = doc.add_paragraph(); rr = q.add_run(line); rr.font.size = Pt(8)
    doc.add_paragraph()
    return t
