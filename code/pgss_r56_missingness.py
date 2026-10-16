# -*- coding: utf-8 -*-
r"""R56（GPT R55 限定只读补齐）：MIMIC-IV / eICU-CRD data v5 冻结合格队列中 23 项冻结预测变量的完整缺失率。只读。

纪律：只读冻结文件；不插补、不删人、不重判表型、不拟合、不算性能；只输出聚合计数，不输出任何个人记录。
数据：19_R38/data_v5（R383 冻结；文件 SHA-256 前 16 位 vs DATA_VERSION.json「v5_SHA」）。
分母：META.eligible = True 的 ICU 住院（GPT 指令：MIMIC-IV 67,000；eICU-CRD 109,104）。
变量：16_R35_FIX/R35_DATA_QC.xlsx「4_最终特征池」（冻结运行代码 pgss_r38_bench.load 读取的同一份清单；
      22 列取自 X，第 23 项「性别(F指示)」按冻结加载器同一表达式由 META.sex 构造）。
字典：itemid / labname / 生理范围 = pgss_r35_extract.py 的 LAB、CHART、E_VITAL 常量（AST 只读解析，不 import——该模块
      导入时要求数据库口令）；单位 = 13_MIMIC_EICU/R34_VARMAP_VERIFIED.xlsx「1_映射与核验」；时间窗 = data_v5 / data_v2 /
      data_v3 DATA_VERSION.json 登记的规则原文。
STOP：某库 X 或 META 指纹不符 → 该库暂停不出数；字典无法确定 → 写「未核实」。
输出：38_R56_ICU预测变量缺失率/R56_ICU_Predictor_Missingness.xlsx 与 R56_QC.json。
"""
import sys, json, hashlib, pathlib, ast
HERE = pathlib.Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
if hasattr(sys.stdout, "reconfigure"): sys.stdout.reconfigure(encoding="utf-8", errors="replace")
import numpy as np, pandas as pd

ROOT = HERE.parent
V5 = ROOT / "19_R38" / "data_v5"
OUT = ROOT / "38_R56_ICU预测变量缺失率"
R55 = pathlib.Path(r"D:\Research\3_对话与交接\GPT交接_PGSS\53\5_GPT_R55交付包（已收到）")
S4DOC = R55 / "02_SUPPLEMENTARY" / "PGSS_Supplementary_v0.2_numbered_review.docx"
POOLX = ROOT / "16_R35_FIX" / "R35_DATA_QC.xlsx"
VARMAP = ROOT / "13_MIMIC_EICU" / "R34_VARMAP_VERIFIED.xlsx"
AUD = ROOT / "19_R38" / "R383_V5_DATA_AUDIT.xlsx"
X51 = ROOT / "34_R51_基线只读提数" / "R51_Table1_Baseline_Source.xlsx"
X49 = ROOT / "31_R49_只读补齐" / "R49_只读补齐_证据表.xlsx"
X48 = ROOT / "30_R48_图稿与补充材料" / "R48_图源与补充材料.xlsx"
EXTRACT = HERE / "pgss_r35_extract.py"
BUILD5 = HERE / "pgss_r383_build_v5.py"
DBS = [("MIMIC-IV", "mimic", 67000), ("eICU-CRD", "eicu", 109104)]      # 分母 = GPT R55 指令
SEX = "性别F"
# 英文名：与 Supplementary Table S4（v0.2）逐行比对，不一致即报错
EN = {"MCV": "MCV", "RDW": "RDW", "SpO2": "SpO2", "呼吸频率": "Respiratory rate", "尿素氮": "Blood urea nitrogen", "平均动脉压": "Mean arterial pressure",
      "年龄": "Age", "心率": "Heart rate", "收缩压": "Systolic blood pressure", "白细胞计数": "White blood cell count", "碳酸氢盐": "Bicarbonate",
      "舒张压": "Diastolic blood pressure", "血小板": "Platelet count", "血氯": "Chloride", "血糖": "Glucose", "血红蛋白": "Hemoglobin", "血细胞比容": "Hematocrit",
      "血肌酐": "Serum creatinine", "血钙": "Calcium", "血钠": "Sodium", "血钾": "Potassium", "阴离子间隙": "Anion gap", SEX: "Female sex"}
# 字典单位（R34 核验映射原文）→ 英文写法；一价离子 mEq/L 与 mmol/L 数值相同，按 Table 1/S5 的 mmol/L 写法并注明原记法
UNIT_EN = {"mg/dL": "mg/dL", "g/dL": "g/dL", "%": "%", "fL": "fL", "K/uL": "×10³/µL", "mEq/L": "mmol/L (recorded as mEq/L)", "bpm": "beats/min",
           "mmHg": "mmHg", "次/分": "breaths/min", "岁": "years"}
ID_COLS = ("stay_id", "subject_id", "hadm_id", "patientunitstayid", "uniquepid", "group_id", "hospital_id", "sid")


def sha16b(b):
    return hashlib.sha256(b).hexdigest()[:16]


def sha16f(p):
    return sha16b(pathlib.Path(p).read_bytes())


def jload(p):
    return json.loads(pathlib.Path(p).read_text(encoding="utf-8"))


# ── 1. 指纹（STOP 规则）──────────────────────────────────────────
def fingerprints():
    man = jload(V5 / "DATA_VERSION.json")
    rows = []
    for db, tag, _ in DBS:
        for t in ("X", "META", "PHENO", "PHENOOK"):
            p = V5 / f"R38_{tag}_{t}.parquet"; got = sha16f(p); reg = man["v5_SHA"][f"{tag}_{t}"]
            use = {"X": "主计数", "META": "合格分母、性别", "PHENOOK": "交叉核验（表型可判定）", "PHENO": "未使用"}[t]
            rows.append(dict(对象=f"data_v5/R38_{tag}_{t}.parquet", 登记来源="data_v5/DATA_VERSION.json v5_SHA", 登记=reg, 实测=got, 一致=got == reg, 本轮用途=use))
    v4 = jload(ROOT / "19_R38" / "data_v4" / "DATA_VERSION.json"); v3 = jload(ROOT / "19_R38" / "data_v3" / "DATA_VERSION.json")
    rows.append(dict(对象="99_代码/pgss_r383_build_v5.py（v5 构建脚本）", 登记来源="data_v5/DATA_VERSION.json 构建脚本SHA", 登记=man["构建脚本SHA"], 实测=sha16f(BUILD5),
                     一致=sha16f(BUILD5) == man["构建脚本SHA"], 本轮用途="时间窗与取值规则出处"))
    rows.append(dict(对象="16_R35_FIX/R35_DATA_QC.xlsx（特征池清单）", 登记来源="data_v3 与 data_v4 DATA_VERSION.json 特征池SHA", 登记=v4["特征池SHA"], 实测=sha16f(POOLX),
                     一致=(sha16f(POOLX) == v4["特征池SHA"] == v3["特征池SHA"]), 本轮用途="23 项变量清单"))
    a5 = pd.read_excel(X48, sheet_name="A5_冻结脚本未改")
    for f, use in (("pgss_r35_extract.py", "itemid / labname / 生理范围字典"), ("pgss_r35_qc.py", "表型可判定规则"), ("pgss_r38_bench.py", "冻结加载器（性别指示构造）")):
        reg = a5[a5.文件 == f]["R47登记SHA16"]; assert len(reg) == 1, f
        got = sha16f(HERE / f)
        rows.append(dict(对象=f"99_代码/{f}", 登记来源="R47 代码包登记（R48 A5_冻结脚本未改）", 登记=reg.iloc[0], 实测=got, 一致=got == reg.iloc[0], 本轮用途=use))
    # eICU 输入在 v3/v4/v5 逐字节不变 ⇒ eICU 时间窗以 v2/v3 登记规则为准
    for vv, m in (("v3", v3["v3_SHA"]), ("v4", v4["v4_SHA"])):
        rows.append(dict(对象=f"eICU X 在 data_{vv} 的登记", 登记来源=f"data_{vv}/DATA_VERSION.json {vv}_SHA eicu_X", 登记=m["eicu_X"], 实测=man["v5_SHA"]["eicu_X"],
                         一致=m["eicu_X"] == man["v5_SHA"]["eicu_X"], 本轮用途="证明 eICU X 自 v3 起未变（时间窗规则沿用 v2/v3 登记）"))
    F = pd.DataFrame(rows)
    # 变异对照：X 文件改一个字节，指纹必须不符
    b = bytearray((V5 / "R38_mimic_X.parquet").read_bytes()); b[len(b) // 2] ^= 0xFF
    mut = sha16b(bytes(b)) != man["v5_SHA"]["mimic_X"]
    c = bytearray(BUILD5.read_bytes()); c[len(c) // 2] ^= 0xFF
    mut2 = sha16b(bytes(c)) != man["构建脚本SHA"]
    ok_db = {db: bool(F[F.对象.isin([f"data_v5/R38_{tag}_X.parquet", f"data_v5/R38_{tag}_META.parquet"])].一致.all()) for db, tag, _ in DBS}
    return F, ok_db, (mut, mut2), man


# ── 2. 变量清单与 S4 对应 ─────────────────────────────────────────
def s4_rows(doc=S4DOC):
    import docx
    d = docx.Document(str(doc))
    caps = [p.text for p in d.paragraphs if p.text.startswith("Supplementary Table S")]
    assert len(caps) == len(d.tables), (len(caps), len(d.tables))
    i = [k for k, c in enumerate(caps) if c.startswith("Supplementary Table S4.")]
    assert len(i) == 1
    t = d.tables[i[0]]; rows = [[c.text for c in r.cells] for r in t.rows]
    assert rows[0][0] == "Predictor", rows[0]
    return caps[i[0]], [r[0] for r in rows[1:]]


def pool_check(internal, s4):
    labels = [EN[v] for v in internal]
    rows = [dict(序号=k + 1, 冻结列名=v, 英文名=labels[k], S4行=(s4[k] if k < len(s4) else None), 一致=(k < len(s4) and labels[k] == s4[k])) for k, v in enumerate(internal)]
    return pd.DataFrame(rows), len(internal) == len(s4) and all(r["一致"] for r in rows)


def pool():
    qc = pd.read_excel(POOLX, sheet_name="4_最终特征池")
    feats = list(qc.特征)
    assert len(feats) == 23 and feats[-1] == "性别(F指示)", feats
    return [{"性别(F指示)": SEX}.get(v, v) for v in feats]


# ── 3. 冻结字典（只读 AST 解析，不 import）──────────────────────────
def frozen_dict():
    src = EXTRACT.read_text(encoding="utf-8"); d = {}
    for node in ast.parse(src).body:
        if isinstance(node, ast.Assign) and len(node.targets) == 1 and isinstance(node.targets[0], ast.Name) and node.targets[0].id in ("LAB", "CHART", "E_VITAL"):
            d[node.targets[0].id] = (ast.literal_eval(node.value), node.lineno)
    assert set(d) == {"LAB", "CHART", "E_VITAL"}, set(d)
    vm = pd.read_excel(VARMAP, sheet_name="1_映射与核验").set_index("变量")
    return d, vm


def dictionary(internal, D, vm, man):
    LAB, l_lab = D["LAB"]; CHART, l_ch = D["CHART"]; EV, l_ev = D["E_VITAL"]
    v2 = jload(ROOT / "19_R38" / "data_v2" / "DATA_VERSION.json"); v3 = jload(ROOT / "19_R38" / "data_v3" / "DATA_VERSION.json")
    w5 = man["v5窗口规则"]
    tw_m = ("Last value before the 24-h landmark: charttime ≥ ICU admission, charttime < landmark and storetime < landmark; latest charttime; "
            "revisions of the same itemid → latest storetime before landmark; several approved itemids at that charttime → arithmetic mean")
    tw_el = ("Last value in the first 24 h of the ICU stay: labresultoffset 0 to <1440 min and labresultrevisedoffset missing or <1440 min; latest result offset; "
             "ties → arithmetic mean (result time; eICU-CRD has no specimen-collection time)")
    tw_ev = "Last value in the first 24 h of the ICU stay: observationoffset 0 to <1440 min; latest offset across the listed tables; ties → arithmetic mean"
    src_m = f"data_v5 DATA_VERSION.json v5窗口规则：{w5['体征体温化验']}；{w5['取值']}"
    src_el = f"data_v2 DATA_VERSION.json v2规则 eICU化验：{v2['v2规则']['eICU化验']}；data_v3 并列规则：{v3['并列规则']}；data_v5 未改动项含「eICU 全部输入」"
    src_ev = f"pgss_r35_extract.py E_VITAL 与 R35 体征查询（observationoffset ∈ [0,1440)）；data_v2 v2规则 eICU体征：{v2['v2规则']['eICU体征']}；data_v3 并列规则：{v3['并列规则']}"
    assert "eICU 全部输入" in man["未改动项"]
    rows = []
    for v in internal:
        r = dict(冻结列名=v, Predictor=EN[v])
        if v in LAB:
            ids, names, (lo, hi) = LAB[v]; u = vm.loc[v, "单位"]
            r.update(Category="Laboratory", Data_type="Continuous", Unit=UNIT_EN[u], 字典单位=u, 单位依据=f"R34_VARMAP_VERIFIED「{vm.loc[v, '换算说明']}」",
                     MIMIC_source=f"hosp.labevents itemid {', '.join(map(str, ids))}", eICU_source=f"lab.labname {', '.join(repr(n) for n in names)}",
                     Range_filter=f"[{lo:g}, {hi:g}] (applied before taking the last value)", MIMIC_window=tw_m, eICU_window=tw_el,
                     字典出处=f"pgss_r35_extract.py LAB（第 {l_lab} 行起）", 时间窗出处_MIMIC=src_m, 时间窗出处_eICU=src_el)
        elif v in CHART:
            ids, (lo, hi) = CHART[v]; u = vm.loc[v, "单位"]
            r.update(Category="Vital sign", Data_type="Continuous", Unit=UNIT_EN[u], 字典单位=u, 单位依据=f"R34_VARMAP_VERIFIED「{vm.loc[v, '换算说明']}」",
                     MIMIC_source=f"icu.chartevents itemid {', '.join(map(str, ids))}", eICU_source=" + ".join(f"{tb}.{col}" for tb, col in EV[v]),
                     Range_filter=f"[{lo:g}, {hi:g}] (applied before taking the last value)", MIMIC_window=tw_m, eICU_window=tw_ev,
                     字典出处=f"pgss_r35_extract.py CHART（第 {l_ch} 行起）、E_VITAL（第 {l_ev} 行起）", 时间窗出处_MIMIC=src_m, 时间窗出处_eICU=src_ev)
        elif v == "年龄":
            r.update(Category="Demographic", Data_type="Continuous", Unit=UNIT_EN[vm.loc[v, "单位"]], 字典单位=vm.loc[v, "单位"], 单位依据=f"R34_VARMAP_VERIFIED「{vm.loc[v, '换算说明']}」",
                     MIMIC_source="hosp.patients anchor_age + (admission year − anchor_year)", eICU_source="patient.age ('> 89' coded 90)",
                     Range_filter="Eligibility requires known age ≥18 years", MIMIC_window="Admission record (time-invariant)", eICU_window="Admission record (time-invariant)",
                     字典出处="pgss_r35_cohort.py build_mimic / build_eicu（年龄表达式）；R51 T4 Age 行", 时间窗出处_MIMIC="不随时间变化", 时间窗出处_eICU="不随时间变化")
        elif v == SEX:
            r.update(Category="Demographic", Data_type="Binary indicator", Unit="1 = female, 0 = male", 字典单位="—", 单位依据="冻结加载器 pgss_r38_bench.load：(meta.sex == 'F').where(meta.sex.isin(['F','M']))",
                     MIMIC_source="hosp.patients gender (F/M)", eICU_source="patient.gender (Female/Male)",
                     Range_filter="Values other than female/male → missing", MIMIC_window="Admission record (time-invariant)", eICU_window="Admission record (time-invariant)",
                     字典出处="pgss_r35_cohort.py（性别映射）；pgss_r38_bench.load", 时间窗出处_MIMIC="不随时间变化", 时间窗出处_eICU="不随时间变化")
        else:
            r.update(Category="未核实", Data_type="未核实", Unit="未核实")
        rows.append(r)
    return pd.DataFrame(rows)


# ── 4. 读取与计数 ────────────────────────────────────────────────
def frame(tag, internal):
    """与冻结加载器 pgss_r38_bench.load 同一构造：X 的 22 列 + META.sex 构造的性别指示；只读。"""
    X = pd.read_parquet(V5 / f"R38_{tag}_X.parquet"); M = pd.read_parquet(V5 / f"R38_{tag}_META.parquet")
    assert X.index.equals(M.index), "X 与 META 索引不一致"
    assert M.eligible.dtype == bool and not M.eligible.isna().any()
    F = X[[v for v in internal if v != SEX]].copy()
    F[SEX] = (M.sex == "F").where(M.sex.isin(["F", "M"]))
    return F[internal], M


def count(F, e, internal, N):
    rows = []
    for v in internal:
        s = pd.to_numeric(F.loc[e, v], errors="coerce").astype(float)
        obs = np.isfinite(s.to_numpy())
        rows.append(dict(冻结列名=v, 分母=N, 观测n=int(obs.sum()), 缺失n=int((~obs).sum()), 无穷值=int(np.isinf(s.to_numpy()).sum()), 原始dtype=str(F[v].dtype)))
    return pd.DataFrame(rows)


def ident_ok(C):
    return bool(((C.观测n + C.缺失n) == C.分母).all())


# ── 5. 交叉核验 ─────────────────────────────────────────────────
def chk_r383(C, db, tag, aud):
    out = []
    for r in C.itertuples(index=False):
        if r.冻结列名 == SEX: continue
        a = aud[(aud.库 == tag) & (aud.变量 == r.冻结列名)]
        assert len(a) == 1, (tag, r.冻结列名)
        ref = int(a.n_available_v5.iloc[0])
        out.append(dict(数据库=db, 冻结列名=r.冻结列名, 本轮观测n=r.观测n, R383_n_available_v5=ref, 一致=r.观测n == ref))
    return out


R51_MAP = {"Female": SEX, "Age": "年龄", "Hemoglobin": "血红蛋白", "Serum creatinine": "血肌酐", "RDW": "RDW", "Platelet count": "血小板", "Chloride": "血氯",
           "eGFR (CKD-EPI 2021)": None, "Albumin": None, "Body mass index": None}


def chk_r51(C, db, t1):
    out = []
    for name, v in R51_MAP.items():
        x = t1[(t1.数据库 == db) & (t1.变量 == name)]; assert len(x) == 1, (db, name)
        x = x.iloc[0]
        if v is None:
            out.append(dict(数据库=db, R51变量=name, 冻结列名="—", R51分母=int(x.分母), R51观测n=int(x.观测n), R51缺失n=int(x.缺失n), 本轮观测n=None, 本轮缺失n=None,
                            一致=None, 说明="不在 23 项预测变量池（只用于表型定义）；不比较"))
            continue
        c = C[C.冻结列名 == v].iloc[0]
        same = int(x.分母) == int(c.分母) and int(x.观测n) == int(c.观测n) and int(x.缺失n) == int(c.缺失n)
        out.append(dict(数据库=db, R51变量=name, 冻结列名=v, R51分母=int(x.分母), R51观测n=int(x.观测n), R51缺失n=int(x.缺失n), 本轮观测n=int(c.观测n), 本轮缺失n=int(c.缺失n),
                        一致=same, 说明="R51 Female 的观测n = 性别已知人数" if v == SEX else ""))
    return out


PH_RULE = {"女性": [SEX], "年龄≥65": ["年龄"], "贫血": ["血红蛋白", SEX], "eGFR降低": ["血肌酐", "年龄", SEX]}


def chk_pheno(F, e, OKf, db, c2):
    """表型可判定 = 定义所需预测变量同时有值（pgss_r35_qc.phenotypes 规则；eGFR 由血肌酐、年龄、性别算出）。
    对冻结 PHENOOK 逐行比对掩码（只报不一致行数），并对 R49 C2 冻结计数。"""
    out = []
    for ph, need in PH_RULE.items():
        joint = np.logical_and.reduce([np.isfinite(pd.to_numeric(F[v], errors="coerce").astype(float).to_numpy()) for v in need])[e.to_numpy()]
        okv = OKf[ph].astype(bool).to_numpy()[e.to_numpy()]
        x = c2[(c2.数据库 == db) & (c2.结局 == "OUT_HOSP_MORT") & (c2.表型 == ph)]; assert len(x) == 1, (db, ph)
        ref = int(x.合格队列.iloc[0]) - int(x.表型不可判定.iloc[0])
        out.append(dict(数据库=db, 表型=ph, 所需预测变量=" + ".join(need), 预测变量同时有值n=int(joint.sum()), 冻结PHENOOK可判定n=int(okv.sum()),
                        逐行不一致数=int((joint != okv).sum()), R49_C2可判定n=ref, 一致=bool(int(joint.sum()) == int(okv.sum()) == ref and (joint == okv).all())))
    return out


def chk_dv(F_all, man, internal):
    """DATA_VERSION.json 聚合统计「有值就诊」是 v5 构建时对 MIMIC 全部 85,242 个首段 ICU 住院（不限合格）的计数。"""
    reg = {r["变量"]: r["有值就诊"] for k in ("MIMIC体征体温聚合", "MIMIC化验聚合") for r in man["聚合统计"][k]}
    out = []
    for v in internal:
        if v not in reg: continue
        n = int(np.isfinite(pd.to_numeric(F_all[v], errors="coerce").astype(float).to_numpy()).sum())
        out.append(dict(数据库="MIMIC-IV", 冻结列名=v, 全部构建住院数=len(F_all), 本轮有值n_全部住院=n, DATA_VERSION有值就诊=int(reg[v]), 一致=n == int(reg[v])))
    return out


def chk_loader(C, db, tag, internal):
    import pgss_r38_bench as BN               # 冻结运行代码的加载器（XGB 主实验经 pgss_xgb_pathaudit 调用同一函数）
    Xf, _, _, meta = BN.load(tag)
    e = meta.eligible.astype(bool)
    cols_same = list(Xf.columns) == internal
    out = []
    for v in internal:
        n = int(np.isfinite(pd.to_numeric(Xf.loc[e, v], errors="coerce").astype(float).to_numpy()).sum())
        c = int(C[C.冻结列名 == v].观测n.iloc[0])
        out.append(dict(数据库=db, 冻结列名=v, 冻结加载器观测n=n, 本轮观测n=c, 一致=n == c))
    return out, cols_same, Xf, meta


R35_NOTE = {("MIMIC-IV", "Laboratory"): "MIMIC 化验时间窗在 v5 改为 charttime ∈ [入 ICU, landmark) 且 storetime < landmark（R35 为 storetime ∈ [入 ICU, landmark)）",
            ("MIMIC-IV", "Vital sign"): "MIMIC 体征时间窗在 v4/v5 改为 charttime ∈ [入 ICU, landmark) 且 storetime < landmark（R35 为 storetime 窗口）",
            ("eICU-CRD", "Laboratory"): "eICU 化验在 v2 加入 labresultrevisedoffset 为空或 < 1440 的条件（24 h 后才修订录入的结果视为 landmark 时不可得）",
            ("eICU-CRD", "Vital sign"): "eICU 体征规则自 R35 未变（v3 只改并列取均值，不影响有无）",
            ("MIMIC-IV", "Demographic"): "未变", ("eICU-CRD", "Demographic"): "未变"}


def chk_r35(C, db, cat, r35):
    out = []
    col = {"MIMIC-IV": "MIMIC-IV", "eICU-CRD": "eICU-CRD"}[db]
    for r in C.itertuples(index=False):
        name = "性别(F指示)" if r.冻结列名 == SEX else r.冻结列名
        x = r35[r35.变量 == name]
        v5 = r.缺失n / r.分母
        old = float(x[col].iloc[0]) if len(x) and pd.notna(x[col].iloc[0]) else None
        out.append(dict(数据库=db, 冻结列名=r.冻结列名, R35选池时缺失率=old, v5缺失率=round(v5, 4), 差_百分点=(None if old is None else round(100 * (round(v5, 4) - old), 2)),
                        低于选池阈值40pct=v5 < 0.40, 说明=("R35 表未列性别（性别不按缺失率选入）" if old is None else R35_NOTE[(db, cat[r.冻结列名])])))
    return out


def safety(dfs):
    bad = []
    for name, d in dfs.items():
        hit = [c for c in d.columns if str(c).lower() in ID_COLS]
        if hit: bad.append((name, hit))
    return bad


# ── 6. 主程序 ───────────────────────────────────────────────────
def fpct(m, N):
    p = 100.0 * m / N
    return "0.0" if m == 0 else ("<0.1" if p < 0.05 else f"{p:.1f}")


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    FP, ok_db, mut_fp, man = fingerprints()
    internal = pool()
    cap, s4 = s4_rows()
    PC, pool_ok = pool_check(internal, s4)
    _, mut_pool = pool_check(internal, list(reversed(s4)))
    D, vm = frozen_dict()
    DIC = dictionary(internal, D, vm, man)
    cat = dict(zip(DIC.冻结列名, DIC.Category))
    aud = pd.read_excel(AUD, sheet_name="2_逐变量v4到v5")
    t1 = pd.read_excel(X51, sheet_name="T1_基线长表"); c2 = pd.read_excel(X49, sheet_name="C2_表型任务")
    r35 = pd.read_excel(POOLX, sheet_name="2_两库缺失率对照")
    QA, MUT = [], {}
    LONG, Q3, Q4, Q5, Q6, Q7, Q8, DEN = [], [], [], [], [], [], [], []
    for db, tag, Ngpt in DBS:
        if not ok_db[db]:
            print(f"STOP：{db} 指纹不符，该库不出数"); continue
        F, M = frame(tag, internal)
        e = M.eligible.astype(bool); N = int(e.sum())
        OKf = pd.read_parquet(V5 / f"R38_{tag}_PHENOOK.parquet").loc[M.index]
        n51 = int(t1[(t1.数据库 == db) & (t1.变量 == "Age")].分母.iloc[0]); n49 = int(c2[(c2.数据库 == db)].合格队列.iloc[0])
        DEN.append(dict(数据库=db, META行数=len(M), 合格N=N, GPT指令分母=Ngpt, R51_T1分母=n51, R49_C2合格队列=n49, 一致=N == Ngpt == n51 == n49,
                        合格标志无缺失=bool(not M.eligible.isna().any()), X与META索引一致=bool(F.index.equals(M.index))))
        MUT[f"分母_{db}_用全部行"] = len(M) != Ngpt
        C = count(F, e, internal, N)
        MUT[f"恒等式_{db}"] = not ident_ok(C.assign(观测n=C.观测n + np.r_[1, np.zeros(len(C) - 1, int)]))
        for r in C.itertuples(index=False):
            d = DIC[DIC.冻结列名 == r.冻结列名].iloc[0]
            LONG.append({"Database": db, "S4 order": internal.index(r.冻结列名) + 1, "Predictor": EN[r.冻结列名], "Frozen column (data v5)": r.冻结列名,
                         "Category": d.Category, "Data type": d.Data_type, "Unit": d.Unit, "Source fields": d.MIMIC_source if db == "MIMIC-IV" else d.eICU_source,
                         "Range filter": d.Range_filter, "Predictor window": d.MIMIC_window if db == "MIMIC-IV" else d.eICU_window,
                         "Eligible ICU stays (N)": r.分母, "Observed n": r.观测n, "Missing n": r.缺失n, "Missing %": round(100.0 * r.缺失n / r.分母, 4),
                         "Observed + missing = N": r.观测n + r.缺失n == r.分母, "Infinite values": r.无穷值})
        Q3 += chk_r383(C, db, tag, aud)
        Cm = count(F.assign(MCV=F.MCV.where(~(F.MCV.notna() & e).cumsum().eq(1))), e, internal, N)      # 变异：第一个合格且有值的 MCV 置缺失
        MUT[f"R383_{db}"] = not all(x["一致"] for x in chk_r383(Cm, db, tag, aud) if x["冻结列名"] == "MCV")
        Q4 += chk_r51(C, db, t1)
        Cm = count(F.assign(血红蛋白=F.血红蛋白.where(~(F.血红蛋白.notna() & e).cumsum().eq(1))), e, internal, N)
        MUT[f"R51_{db}"] = not all(x["一致"] for x in chk_r51(Cm, db, t1) if x["冻结列名"] == "血红蛋白")
        Q5 += chk_pheno(F, e, OKf, db, c2)
        k = F.index[(F.血红蛋白.notna() & F[SEX].notna() & e).to_numpy()][0]
        Fm = F.copy(); Fm[SEX] = Fm[SEX].astype(object); Fm.loc[k, SEX] = np.nan
        MUT[f"表型_{db}"] = not all(x["一致"] for x in chk_pheno(Fm, e, OKf, db, c2) if x["表型"] == "贫血")
        if tag == "mimic":
            Q6 += chk_dv(F, man, internal)
            MUT["DATA_VERSION_MIMIC"] = not all(x["一致"] for x in chk_dv(F.iloc[1:], man, internal))
        lo, cols_same, Xf, meta = chk_loader(C, db, tag, internal)
        Q7 += lo
        QA.append(dict(检查=f"冻结加载器列顺序 = 23 项清单（{db}）", 通过=cols_same, 变异对照检出="不适用", 说明=f"{len(Xf.columns)} 列"))
        sw = C.copy(); sw.loc[sw.冻结列名 == "SpO2", "观测n"], sw.loc[sw.冻结列名 == "年龄", "观测n"] = int(C[C.冻结列名 == "年龄"].观测n.iloc[0]), int(C[C.冻结列名 == "SpO2"].观测n.iloc[0])
        lo_m, _, _, _ = chk_loader(sw, db, tag, internal)
        MUT[f"加载器_{db}"] = not all(x["一致"] for x in lo_m)
        Q8 += chk_r35(C, db, cat, r35)
    L = pd.DataFrame(LONG)
    # 打印版（宽表，23 行）
    pr = []
    for v in internal:
        r = {"Predictor": EN[v], "Data type": DIC[DIC.冻结列名 == v].Data_type.iloc[0], "Unit": DIC[DIC.冻结列名 == v].Unit.iloc[0]}
        for db, _, _ in DBS:
            x = L[(L.Database == db) & (L["Frozen column (data v5)"] == v)]
            if len(x):
                x = x.iloc[0]
                r[f"{db} observed n"] = f"{x['Observed n']:,}"; r[f"{db} missing n (%)"] = f"{x['Missing n']:,} ({fpct(x['Missing n'], x['Eligible ICU stays (N)'])})"
            else:
                r[f"{db} observed n"] = r[f"{db} missing n (%)"] = "PAUSED"
        pr.append(r)
    P = pd.DataFrame(pr)
    D3, D4, D5, D6, D7, D8, DN = (pd.DataFrame(x) for x in (Q3, Q4, Q5, Q6, Q7, Q8, DEN))
    out_dfs = {"1_S5A_long": L, "2_S5A_print": P, "3_Dictionary": DIC}
    sf = safety(out_dfs); sf_m = safety({"变异": L.assign(stay_id=0)})
    QA = [dict(检查="data v5 指纹：两库 X/META/PHENO/PHENOOK 8 个文件 = DATA_VERSION.json v5_SHA", 通过=bool(FP.iloc[:8].一致.all()), 变异对照检出=mut_fp[0], 说明=f"{int(FP.iloc[:8].一致.sum())}/8"),
          dict(检查="构建脚本、特征池清单、字典脚本、加载器、eICU 输入链的登记指纹", 通过=bool(FP.iloc[8:].一致.all()), 变异对照检出=mut_fp[1], 说明=f"{int(FP.iloc[8:].一致.sum())}/{len(FP) - 8}"),
          dict(检查="23 项清单与 Supplementary Table S4（v0.2）逐行同名同序", 通过=pool_ok, 变异对照检出=not mut_pool, 说明=f"{int(PC.一致.sum())}/{len(PC)}；{cap[:60]}"),
          dict(检查="分母 = GPT 指令 = R51 T1 = R49 C2（合格标志无缺失、X/META 索引一致）", 通过=bool(DN.一致.all() and DN.合格标志无缺失.all() and DN.X与META索引一致.all()),
               变异对照检出=all(v for k, v in MUT.items() if k.startswith("分母")), 说明="；".join(f"{r.数据库} {r.合格N:,}" for r in DN.itertuples())),
          dict(检查="46 行：观测n + 缺失n = 分母；无无穷值", 通过=bool(L["Observed + missing = N"].all() and (L["Infinite values"] == 0).all() and len(L) == 46),
               变异对照检出=all(v for k, v in MUT.items() if k.startswith("恒等式")), 说明=f"{len(L)} 行"),
          dict(检查="对 R383 v5 构建审计（2_逐变量v4到v5 n_available_v5，合格分母）逐值相同", 通过=bool(D3.一致.all()), 变异对照检出=all(v for k, v in MUT.items() if k.startswith("R383")), 说明=f"{int(D3.一致.sum())}/{len(D3)}"),
          dict(检查="对 R51 Table 1 基线长表（与预测变量池重合的 7 项）逐值相同", 通过=bool(D4.一致.dropna().all()), 变异对照检出=all(v for k, v in MUT.items() if k.startswith("R51")),
               说明=f"{int(D4.一致.dropna().sum())}/{int(D4.一致.notna().sum())}；另 {int(D4.一致.isna().sum())} 行（eGFR/白蛋白/BMI）不在池中不比较"),
          dict(检查="表型可判定 = 定义所需预测变量同时有值：对冻结 PHENOOK 逐行一致、对 R49 C2 计数一致", 通过=bool(D5.一致.all()), 变异对照检出=all(v for k, v in MUT.items() if k.startswith("表型")), 说明=f"{int(D5.一致.sum())}/{len(D5)}"),
          dict(检查="对 DATA_VERSION.json 聚合统计「有值就诊」（MIMIC 全部 85,242 构建住院）逐值相同", 通过=bool(D6.一致.all()), 变异对照检出=MUT.get("DATA_VERSION_MIMIC"), 说明=f"{int(D6.一致.sum())}/{len(D6)}"),
          dict(检查="冻结运行代码加载器 pgss_r38_bench.load 的同一矩阵计数相同", 通过=bool(D7.一致.all()), 变异对照检出=all(v for k, v in MUT.items() if k.startswith("加载器")), 说明=f"{int(D7.一致.sum())}/{len(D7)}"),
          dict(检查="输出表不含个人标识列（只有聚合行）", 通过=not sf, 变异对照检出=bool(sf_m), 说明=f"{sum(len(d) for d in out_dfs.values())} 行聚合"),
          ] + QA
    QA = pd.DataFrame(QA)
    readme = pd.DataFrame({"Item": ["Title", "Status", "Denominators", "Missing definition", "Not done", "Frozen inputs", "Sheets"],
                           "Content": ["Supplementary Table S5A (candidate number; final numbering to be decided by GPT). Completeness of the 23 frozen predictors in the eligible MIMIC-IV and eICU-CRD cohorts (data v5).",
                                       "Read-only count of the frozen predictor matrix; no fitting, imputation, exclusion or phenotype reassignment. Current S1–S44 numbering, main text and Supplementary Word were not modified.",
                                       "; ".join(f"{r.数据库}: {r.合格N:,} eligible ICU stays" for r in DN.itertuples()),
                                       "A value is missing when no measurement passing the frozen range filter was available in the predictor window (female sex: sex not recorded as female or male). Counts describe the frozen matrix before any model-specific handling.",
                                       "No new extraction, no unit re-scan, no NWICU/NHANES additions.",
                                       "19_R38/data_v5 X and META parquet files (SHA-256 prefixes match DATA_VERSION.json); variable list R35_DATA_QC.xlsx 4_最终特征池 (registered feature-pool SHA).",
                                       "1_S5A_long (46 rows), 2_S5A_print (23 rows), 3_Dictionary (23 rows), QA_* (quality checks), QA_0_summary."]})
    with pd.ExcelWriter(OUT / "R56_ICU_Predictor_Missingness.xlsx", engine="openpyxl") as xw:
        readme.to_excel(xw, sheet_name="0_README", index=False)
        L.to_excel(xw, sheet_name="1_S5A_long", index=False); P.to_excel(xw, sheet_name="2_S5A_print", index=False); DIC.to_excel(xw, sheet_name="3_Dictionary", index=False)
        QA.to_excel(xw, sheet_name="QA_0_summary", index=False); FP.to_excel(xw, sheet_name="QA_1_fingerprints", index=False); PC.to_excel(xw, sheet_name="QA_2_S4_mapping", index=False)
        DN.to_excel(xw, sheet_name="QA_3_denominators", index=False); D3.to_excel(xw, sheet_name="QA_4_vs_R383_audit", index=False); D4.to_excel(xw, sheet_name="QA_5_vs_R51_Table1", index=False)
        D5.to_excel(xw, sheet_name="QA_6_phenotype_ascert", index=False); D6.to_excel(xw, sheet_name="QA_7_vs_DATA_VERSION", index=False); D7.to_excel(xw, sheet_name="QA_8_frozen_loader", index=False)
        D8.to_excel(xw, sheet_name="QA_9_vs_R35_pool_QC", index=False)
    top = L.sort_values("Missing %", ascending=False).groupby("Database").head(3)
    res = dict(QA=f"{int(QA.通过.sum())}/{len(QA)}", 通过=bool(QA.通过.all()), 变异对照=MUT, 停止的库=[db for db, ok in ok_db.items() if not ok],
               分母={r.数据库: r.合格N for r in DN.itertuples()}, 行数=len(L),
               各库=[dict(数据库=db, 缺失为0项=int(((L.Database == db) & (L["Missing n"] == 0)).sum()), 缺失率最高=[dict(变量=r.Predictor, 缺失n=int(r["Missing n"]), 缺失pct=round(r["Missing %"], 2)) for _, r in top[top.Database == db].iterrows()],
                        中位缺失pct=round(float(L[L.Database == db]["Missing %"].median()), 2), 最大缺失pct=round(float(L[L.Database == db]["Missing %"].max()), 2)) for db, _, _ in DBS],
               QA明细=QA.to_dict("records"), 指纹=FP.to_dict("records"), S4标题=cap, R35最大差=D8.dropna(subset=["差_百分点"]).groupby("数据库").差_百分点.agg(["min", "max"]).reset_index().to_dict("records"),
               R35说明=sorted(set(D8.说明)))
    (OUT / "R56_QC.json").write_text(json.dumps(res, ensure_ascii=False, indent=1, default=str), encoding="utf-8")
    print(QA[["检查", "通过", "变异对照检出", "说明"]].to_string())
    print("QA", res["QA"], "；变异对照", sum(bool(v) for v in MUT.values()), "/", len(MUT))
    if not res["通过"] or not all(MUT.values()):
        raise SystemExit("R56 QA 未全过")


if __name__ == "__main__":
    main()
