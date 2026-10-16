# -*- coding: utf-8 -*-
r"""R54：PGSS Supplementary Information v0.1 构建 + 第一阶段只读核验 + QA。只读冻结表，不拟合、不计算新统计量。
输出（37_R54_补充材料v0.1/）：PGSS_Supplementary_v0.1.docx、PGSS_Supplementary_v0.1_DataTables.xlsx、R54_SUPP_INDEX.xlsx、R54_QA.json、R54_PHASE1.json
"""
import sys, re, json, pathlib, difflib, hashlib, time
HERE = pathlib.Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
if hasattr(sys.stdout, "reconfigure"): sys.stdout.reconfigure(encoding="utf-8", errors="replace")
import numpy as np, pandas as pd
import pgss_r54_supp_common as C
import pgss_r54_supp_tables as S

OUT = C.ROOT / "37_R54_补充材料v0.1"
PK = pathlib.Path(r"D:\Research\3_对话与交接\GPT交接_PGSS\53\3_GPT_R54交付包（已收到）")
XROOT = pathlib.Path(r"D:\Research\3_对话与交接\GPT交接_PGSS")
TOK = re.compile(r"[−+\-]?\d[\d,]*(?:\.\d+)?")
FORBID = {"SEQN", "unit_id", "uid", "stay_id", "subject_id", "hadm_id", "group_id", "patientunitstayid", "uniquepid", "sid"}


# ── 第一阶段：v0.5a 差异与 S1 核验 ───────────────────────────────
def phase1():
    import docx
    from PIL import Image
    def flat(f):
        d = docx.Document(f)
        out = [("P", p.text.strip()) for p in d.paragraphs if p.text.strip()]
        for ti, t in enumerate(d.tables):
            for ri, r in enumerate(t.rows):
                for ci, c in enumerate(r.cells):
                    out.append((f"T{ti + 1}r{ri}c{ci}", c.text.strip()))
        return out
    arch = PK / "01_MANUSCRIPT" / "ARCHIVE_PGSS_Manuscript_v0.5_R53_integrated_review.docx"
    r53 = XROOT / "53" / "1_GPT_R53交付包（已收到）" / "PGSS_R53_交付包" / "01_MANUSCRIPT" / "PGSS_Manuscript_v0.5_R53_integrated_review.docx"
    a, b = flat(arch), flat(PK / "01_MANUSCRIPT" / "PGSS_Manuscript_v0.5a_R54_two_edits.docx")
    sm = difflib.SequenceMatcher(a=[x[1] for x in a], b=[x[1] for x in b], autojunk=False)
    ch = []
    for op, i1, i2, j1, j2 in sm.get_opcodes():
        if op != "equal":
            ch.append(dict(op=op, old=[f"{x[0]}: {x[1]}" for x in a[i1:i2]], new=[f"{x[0]}: {x[1]}" for x in b[j1:j2]]))
    expected_new = ["Eligible eICU-CRD stays came from 207 hospitals; the other frozen data files do not encode a hospital identifier.",
                    "Low-albumin B track showed +3.23% interval-supported gain.", "Clinical FCT (MIMIC-IV/eICU-CRD)"]
    newtxt = " ".join(" ".join(c["new"]) for c in ch)
    nums_a = sorted(set(TOK.findall(" ".join(x[1] for x in a)))); nums_b = sorted(set(TOK.findall(" ".join(x[1] for x in b))))
    s1 = json.loads((PK / "06_QA" / "R54_S1_verified_frozen_panels.json").read_text(encoding="utf-8"))
    flow = {p["title"].split()[1] if p["title"].split()[1] != "NHANES" else "NHANES": [st[1] for st in p["steps"]] for p in s1["panels"]}
    ref = {"MIMIC-IV": [85242, 67222, 67222, 67000], "eICU-CRD": [158442, 110431, 110161, 109104], "NWICU": [28612, 26764, 20057, 20057, 20011, 19965],
           "NHANES": [51623, 26246, 24693, 24666, 23513]}   # R51/R53 已核实的 58 个数字中的纳排链（来源：R49 C1、R48 S2、R44 COHORT）
    c1 = C.sheet("R49", "C1_纳排流程"); s2 = C.sheet("R48", "S2_E1_NWICU_队列流程"); s9 = C.sheet("R48", "S9_A2_纳排流程")
    src = {"MIMIC-IV": c1[(c1.数据库 == "MIMIC-IV") & c1.步骤.notna()].sort_values("步骤").住院单位数.astype(int).tolist(),
           "eICU-CRD": c1[(c1.数据库 == "eICU-CRD") & c1.步骤.notna()].sort_values("步骤").住院单位数.astype(int).tolist(),
           "NWICU": s2.人次.astype(int).tolist(), "NHANES": s9.剩余.astype(int).tolist()}
    details = " ".join(p["details"] for p in s1["panels"])
    outnums = {"53,539", "7,193", "27,068", "94,920", "207", "9,569", "39,609", "19,965", "17,812", "8,146", "19,765", "17,626", "2,100", "23,513", "3,813", "23,512", "1,243"}
    png = Image.open(PK / "02_SUPPLEMENTARY" / "PGSS_R54_SuppFig_S1_Cohort_flow.png"); tif = Image.open(PK / "02_SUPPLEMENTARY" / "PGSS_R54_SuppFig_S1_Cohort_flow.tif")
    A = np.asarray(png.convert("RGB")); B = np.asarray(tif.convert("RGB"))
    cap = (PK / "02_SUPPLEMENTARY" / "PGSS_R54_SuppFig_S1_caption.txt").read_text(encoding="utf-8").splitlines()[0]
    rec = dict(v05_archive_equals_R53=arch.read_bytes() == r53.read_bytes(), 比对元素=len(a), 改动处=len(ch), 改动=ch,
               三处修改均为预期文字=all(e in newtxt for e in expected_new), 数字集合不变=nums_a == nums_b,
               S1_节点与冻结来源一致={k: flow.get(k) == src[k] == ref[k] for k in ref}, S1_结局框数字齐全=all(n in details for n in outnums),
               S1_PNG与TIFF像素相同=bool(A.shape == B.shape and np.array_equal(A, B)), S1_TIFF_dpi=list(tif.info.get("dpi", ())), S1_TIFF_压缩=tif.info.get("compression"),
               S1_图题与图注一致=bool("Cohort assembly and outcome eligibility by data source" in cap and s1.get("version", "").startswith("R54")))
    (OUT / "R54_PHASE1.json").write_text(json.dumps(rec, ensure_ascii=False, indent=1, default=str), encoding="utf-8")
    return rec


# ── 来源矩阵 ─────────────────────────────────────────────────────
def source_matrix(tables):
    s0 = pd.read_csv(PK / "03_EDITORIAL" / "R54_R48_S0_原始状态索引.csv", encoding="utf-8-sig")
    used = {}
    for t in tables:
        for src, sh, _ in t["src"]:
            used.setdefault((src, sh), []).append(t["id"])
        for name, _, lab in t["long"]:
            for part in lab.split(" + "):
                used.setdefault(("LABEL", part.strip().split(" (")[0]), []).append(f"data:{name}")
    r48 = pd.ExcelFile(C.SRC["R48"]).sheet_names; r47 = pd.ExcelFile(C.SRC["R47"]).sheet_names
    rows = []
    for r in s0.itertuples():
        sid = r.建议编号 if isinstance(r.建议编号, str) else "—"
        status = r.处理
        if status == "R47包内":
            avail = "R47 local workbook" if r.工作表 in r47 else ("R47 code package (t2_records)" if "t2_records" in str(r.工作表) else "not found")
            key_print = [("R47", r.工作表)]; key_long = ("LABEL", f"R47 {r.工作表.split('_')[0]}")
        elif status in ("接回", "本轮新增"):
            sh = r.本包表名 if isinstance(r.本包表名, str) else r.工作表
            avail = "R48 workbook" if sh in r48 else ("R48 figure-source sheets" if "图源" in str(r.工作表) else "R48 package (non-table)")
            key_print = [("R48", sh)]; key_long = ("LABEL", f"R48 {sh.split('_')[0]}")
        elif status == "仅索引":
            p = XROOT / str(r.交付副本) if isinstance(r.交付副本, str) else None
            ok = bool(p and p.exists() and r.工作表 in pd.ExcelFile(p).sheet_names)
            avail = f"original workbook present locally ({r.交付副本})" if ok else "not physically re-attached; original not located"
            key_print = []; key_long = None
        elif status == "已被取代":
            avail = "superseded — not used"; key_print = []; key_long = None
        else:
            avail = "gap (closed by R49)" if "纳排" in str(r.工作表) or "辅助指标" in str(r.工作表) else "gap"
            key_print = []; key_long = None
        uses = [i for k in key_print for i in used.get(k, [])] + (used.get(key_long, []) if key_long else [])
        if sid == "S85":
            decision = "INTERNAL ONLY (verbatim literature excerpts; replace by short citations in the manuscript)"
        elif sid == "S78":
            decision = "SEPARATE RESOURCE (synthetic per-replicate records; distribute with code, not printed)"
        elif status == "已被取代":
            decision = f"EXCLUDED: superseded by {r.取代者 if isinstance(r.取代者, str) else 'later version'}"
        elif status == "仅索引":
            decision = "NOT INCLUDED in v0.1 (index-only)" if not uses else "used"
        elif uses:
            decision = "USED: " + ", ".join(sorted(set(uses)))
        elif status == "缺口":
            decision = "Gap closed by R49 (see S91–S95 rows)"
        else:
            decision = "NOT USED in v0.1 (available; see remarks)"
        if sid in SPECIAL_USE:
            decision = SPECIAL_USE[sid]
        elif decision.startswith("NOT USED") and sid in NOTUSED_REASON:
            decision = "NOT PRINTED: " + NOTUSED_REASON[sid]
        rows.append(dict(internal_id=sid, R48_group=r.组, source=r.来源, worksheet=r.工作表, original_status=status, rows=r.行数, availability=avail, v0_1_decision=decision))
    for sid, sh, desc in (("S91", "C1_纳排流程", "MIMIC-IV/eICU-CRD cohort flow"), ("S92", "C2_表型任务", "phenotype-task sizes"), ("S93", "C3_定义与landmark", "definitions and landmark"),
                          ("S94", "D1_各模型跨seed汇总", "NWICU/NHANES cross-seed auxiliary summaries"), ("S95", "D3_配对差汇总_R43格式", "four-source R43-format summaries")):
        uses = used.get(("R49", sh), []) + used.get(("LABEL", f"R49 {sh.split('_')[0]}"), [])
        rows.append(dict(internal_id=sid, R48_group="R49 increment", source="R49 evidence workbook", worksheet=sh, original_status="R49 new", rows=len(C.sheet("R49", sh)),
                         availability="R49 workbook", v0_1_decision=SPECIAL_USE.get(sid) or (("USED: " + ", ".join(sorted(set(uses)))) if uses else "NOT USED in v0.1")))
    for src, sh, desc in (("R51", "T1_基线长表", "baseline long table"), ("R51", "T3_表型可判定覆盖率", "phenotype ascertainment"), ("R51", "T4_单位核实", "unit verification"),
                          ("R52", "8_E1主对照", "creatinine sensitivity E1"), ("R52", "16_L3预定判读", "creatinine sensitivity L3")):
        uses = used.get((src, sh), [])
        rows.append(dict(internal_id="—", R48_group=f"{src} increment", source=f"{src} workbook", worksheet=sh, original_status=f"{src} new", rows=len(C.sheet(src, sh)),
                         availability=f"{src} workbook", v0_1_decision=("USED: " + ", ".join(sorted(set(uses)))) if uses else "NOT USED in v0.1"))
    return pd.DataFrame(rows)


SPECIAL_USE = {"S2": "USED: Supplementary Figure S1 panel C (verified in Phase 1)", "S9": "USED: Supplementary Figure S1 panel D (verified in Phase 1)",
               "S91": "USED: Supplementary Figure S1 panels A–B (verified in Phase 1); reasons in SM1.1",
               "S84": "USED: rendered as Supplementary Methods 2 text", "S12": "USED: QA cross-checks of ST7.1 and ST8.1 (pooled three-level table; per-source rows printed)"}
NOTUSED_REASON = {"S13": "replaced by ST5.3 if listed; otherwise counts replicated in ST6.2/ST6.3",
                  "S36": "both rows are contained in ST8.4 (R48 S41, 12 rows)",
                  "S45": "identical MIMIC-IV + eICU-CRD rows are included in ST9.1 via R49 D3 (labelled = R43 B3)",
                  "S52": "identical MIMIC-IV + eICU-CRD rows are included in ST9.2 via R49 D4 (labelled = R43 B4)",
                  "S86": "reproduction entry points belong to the code package README, not the printed supplement",
                  "S87": "figure-legend checklist for main Figures 1–6 (manuscript legends), not a supplementary table",
                  "S88": "source data of main Figures 1–6; to be released as figure source data, not printed",
                  "S90": "local file paths; internal reproduction aid, not distributable"}


def sup_used(tables, sup_sheets):
    return sorted({h for t in tables for s, h, _ in t["src"] if h in sup_sheets and s not in ("R47", "R48", "R49", "R51", "R52")} |
                  {h for t in tables for s, h, _ in t["src"] if s in ("R42", "R43", "R44", "R45", "R46")})


# ── 数字溯源 QA ──────────────────────────────────────────────────
def norm_tok(t):
    return t.replace("−", "-").replace("+", "").replace(",", "")


def src_strings(vals, dps):
    out = set()
    for d in dps:
        for v in vals:
            for x in (v, 100 * v):
                s = f"{x:.{d}f}"
                out.add(s); out.add(s.lstrip("-") if s.strip("-0.") == "" else s)
    return out


def trace(t):
    df = t["df"]; bad = []
    if not t["num_cols"]:
        return 0, bad
    toks = [(c, i, x) for c in t["num_cols"] for i, v in enumerate(df[c].astype(str)) for x in TOK.findall(v)]
    dps = sorted({len(norm_tok(x).split(".")[1]) if "." in x else 0 for _, _, x in toks})
    allowed = src_strings(t["src_vals"], dps)
    for c, i, x in toks:
        n = norm_tok(x)
        if n not in allowed and n.lstrip("-") not in allowed and ("-" + n) not in allowed:
            bad.append(f"{c}[{i}]={x}")
    return len(toks), bad


def cross_checks(tables):
    """用另一份冻结工作簿独立核对三张关键表（不同来源 = 非自比）。"""
    tb = {t["id"]: t for t in tables}
    out = []
    f2 = C.sheet("R47", "F2_ICU复现")       # R47 图 2 图源（ICU 36 任务 E1 两轨等权）
    e = tb["ST5.1"]["df"]
    ok = 0
    for r in f2.itertuples():
        row = e[(e.Database == C.DB[r.库]) & (e.Outcome == C.OC[r.结局]) & (e.Phenotype == C.PH[r.表型])]
        ok += len(row) == 1 and row.iloc[0]["Relative improvement (%)"] == C.fpct(r.观察相对改善, 2, True) and row.iloc[0]["Pointwise 95% interval (%)"] == C.fci(r.相对_P2_5, r.相对_P97_5)
    out.append(dict(检查="ST5.1 每行 = R47 F2_ICU复现（图 2 图源）", 通过=ok == len(f2) == len(e), 说明=f"{ok}/{len(f2)}"))
    b2 = C.sheet("R47", "B2_三层总结果表")
    l2 = b2[b2.层.str.startswith("L2") & (b2.对照 == "E1")]
    s = tb["ST7.1"]["df"]; ok = 0
    for r in l2.itertuples():
        row = s[(s.Outcome == C.OC[r.结局]) & (s.Phenotype == C.PH[r.表型]) & (s.Scale == C.SC[r.尺度])]
        ok += len(row) == 1 and row.iloc[0]["Relative improvement (%)"] == C.fpct(r.观察相对改善, 2, True) and row.iloc[0]["m = 1% judgment"] == C.TS[r.三态_m1]
    out.append(dict(检查="ST7.1 每行 = R47 B2_三层总结果表（L2 E1，全部尺度）", 通过=ok == len(l2) == len(s), 说明=f"{ok}/{len(l2)}"))
    l3 = b2[b2.层.str.startswith("L3")]
    j = tb["ST8.1"]["df"]; ok = 0
    for r in l3.itertuples():
        row = j[(j.Contrast == r.对照) & (j.Outcome == C.OC[r.结局]) & (j.Phenotype == C.PH[r.表型]) & (j["Judgment scale"] == C.SC[r.尺度])]
        ok += len(row) == 1 and row.iloc[0]["Relative improvement (%)"] == C.fpct(r.观察相对改善, 2, True)
    n_judged = int(((l3.对照 == "C1") & (l3.尺度 == "pop")).sum() + ((l3.对照 == "D1") & (l3.尺度 == "B")).sum())
    out.append(dict(检查="ST8.1 每行 = R47 B2_三层总结果表（L3 预定判读尺度）", 通过=ok == n_judged == len(j), 说明=f"{ok}/{n_judged}"))
    return out


def qa(tables, longs, idx, smx, docx_path):
    R = []
    def add(name, ok, mut, det=""):
        R.append(dict(检查=name, 通过=bool(ok), 变异对照检出=(bool(mut) if mut is not None else "不适用"), 说明=str(det)[:500]))
    ntok, badall = 0, []
    for t in tables:
        n, bad = trace(t); ntok += n; badall += [f"{t['id']}:{b}" for b in bad]
    # 变异：ST5.1 第一行相对改善末位 +1
    t0 = next(t for t in tables if t["id"] == "ST5.1"); m = dict(t0); m["df"] = t0["df"].copy()
    v = m["df"].loc[0, "Relative improvement (%)"]; last = int(v[-1]); m["df"].loc[0, "Relative improvement (%)"] = v[:-1] + str((last + 1) % 10)
    add("1 印刷表每个数字都能在声明的来源列中找到（按小数位比较原值或百分数）", not badall, bool(trace(m)[1]), f"{ntok} 个数字；未找到 {badall[:8]}")
    for c in cross_checks(tables):
        add("2 " + c["检查"], c["通过"], None, c["说明"])
    # 交叉核对的变异：改 ST5.1 一格后重跑
    mt = [dict(t) for t in tables]; i5 = next(i for i, t in enumerate(mt) if t["id"] == "ST5.1"); mt[i5] = m
    add("2b 交叉核对对改动灵敏（ST5.1 改一格）", True, not cross_checks(mt)[0]["通过"])
    sup = smx[smx.original_status == "已被取代"]; sset = set(sup.worksheet)
    fake = dict(src=[("R45", "X10_NWICU_Brier全部对照", 528)])
    add("3 没有直接引用任何已被取代（superseded）或旧轮次的来源表（只从 R47/R48/R49/R51/R52 冻结工作簿取数）", not sup_used(tables, sset),
        bool(sup_used(tables + [fake], sset)), f"已被取代 {len(sup)} 条全部排除")
    cj = [(n, C.residual_cjk(d)) for n, d in [(t["id"], t["df"]) for t in tables] + [(k, v[0]) for k, v in longs.items()]]
    cj = [(n, r) for n, r in cj if r[0] or r[1]]
    probe = pd.DataFrame({"x": ["贫血"]})
    add("4 印刷表与长表无残留中文", not cj, bool(C.residual_cjk(probe)[1]), cj[:5])
    priv = [k for k, v in longs.items() if set(map(str, v[0].columns)) & FORBID]
    add("5 长表不含个人标识列", not priv, bool(set(map(str, pd.DataFrame({"SEQN": [1]}).columns)) & FORBID), f"{len(longs)} 张长表")
    full = idx[idx.fidelity == "full"]
    mism = [r.item for r in full.itertuples() if pd.notna(r.source_rows_primary) and r.printed_rows != r.source_rows_primary and r.item not in FULL_EXEMPT]
    idxm = full.copy(); k = idxm.index[~idxm.item.isin(FULL_EXEMPT)][0]; idxm.loc[k, "printed_rows"] += 1
    mism_m = [r.item for r in idxm.itertuples() if pd.notna(r.source_rows_primary) and r.printed_rows != r.source_rows_primary and r.item not in FULL_EXEMPT]
    add("6 标为 full 的表印刷行数 = 主来源行数（例外项已在备注说明）", not mism, bool(mism_m), mism)
    import docx
    d = docx.Document(docx_path)
    caps = [p.text for p in d.paragraphs if p.text.startswith("Supplementary Table ") and p.runs and p.runs[0].bold]   # 只数加粗表题（目录行不加粗）
    add("7 Word 中表格数与加粗表题数 = 定义数", len(d.tables) == len(tables) == len(caps), None, f"tables {len(d.tables)} / captions {len(caps)} / specs {len(tables)}")
    def cells_ok(tabs):
        return all(d.tables[i].rows[1].cells[0].text == str(t["df"].iloc[0, 0]) and d.tables[i].rows[-1].cells[-1].text == str(t["df"].iloc[-1, -1]) for i, t in enumerate(tabs))
    mt8 = [dict(t) for t in tables]; mt8[3] = dict(mt8[3]); mt8[3]["df"] = mt8[3]["df"].copy(); mt8[3]["df"].iloc[0, 0] = str(mt8[3]["df"].iloc[0, 0]) + "x"
    add("8 Word 表格内容与构建结果一致（每表首格与末格）", cells_ok(tables), not cells_ok(mt8), "")
    return pd.DataFrame(R)


# 印刷行数 ≠ 主来源行数但仍为全量的表（原因写在索引 remarks）
FULL_EXEMPT = {"ST5.4", "ST1.1", "ST2.1", "ST2.2", "ST3.1", "ST3.2", "ST4.1", "ST4.2", "ST5.2", "ST6.1", "ST8.1", "ST9.2", "ST11.1", "ST14.3", "ST15.2", "SM1.1"}


def build_docx(tables, path):
    import docx
    from docx.shared import Pt, Cm
    from docx.enum.section import WD_ORIENT
    d = docx.Document()
    st = d.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(10)
    sec = d.sections[0]; sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(1.6))
    h = d.add_paragraph(); r = h.add_run("Supplementary Information (draft v0.1 for review)"); r.bold = True; r.font.size = Pt(14)
    d.add_paragraph("Phenotype-specific versus phenotype-aware pooled learning: a multicohort benchmark and simulation study")
    d.add_paragraph("Draft compiled read-only from frozen summary workbooks. Item labels (SM1, ST1.1, …) are provisional thematic codes and are not final journal numbering; "
                    "internal source identifiers are listed in the accompanying index. Long tables (all contrasts, per-seed metrics, simulation details) are supplied as a machine-readable workbook.")
    p = d.add_paragraph(); rr = p.add_run("Contents"); rr.bold = True
    for line in ["Supplementary Methods 1 and Supplementary Figure S1"] + [t["title"] for t in tables] + ["Supplementary Methods 2"]:
        q = d.add_paragraph(line); q.paragraph_format.space_after = Pt(0)
        for run in q.runs:
            run.font.size = Pt(8)
    d.add_page_break()
    def para(txt, bold=False):
        q = d.add_paragraph(); run = q.add_run(txt); run.bold = bold
    para(S.SM1_TEXT[0], True)
    for txt in S.SM1_TEXT[1:]:
        para(txt)
    C.add_table(d, tables[0]["title"], tables[0]["df"], tables[0]["legend"])
    para("Supplementary Figure S1. Cohort assembly and outcome eligibility by data source.", True)
    d.add_picture(str(PK / "02_SUPPLEMENTARY" / "PGSS_R54_SuppFig_S1_Cohort_flow.png"), width=Cm(22))
    cap = (PK / "02_SUPPLEMENTARY" / "PGSS_R54_SuppFig_S1_caption.txt").read_text(encoding="utf-8").splitlines()
    for line in cap[1:]:
        if line.strip():
            q = d.add_paragraph(); run = q.add_run(line.strip()); run.font.size = Pt(8)
    d.add_page_break()
    group = None
    for t in tables[1:]:
        if t["group"] == "SM2" and group != "SM2":
            d.add_page_break(); para(S.SM2_TEXT[0], True)
            for txt in S.SM2_TEXT[1:]:
                para(txt)
        group = t["group"]
        C.add_table(d, t["title"], t["df"], t["legend"], font_pt=7 if len(t["df"].columns) > 8 else 7.5)
    d.save(path)


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    t0 = time.time()
    p1 = phase1()
    tables = S.all_tables()
    longs, seen = {}, {}
    for t in tables:
        for name, df, lab in t["long"]:
            assert name not in longs and len(name) <= 31, name
            longs[name] = (df, lab, t["id"])
    docx_path = OUT / "PGSS_Supplementary_v0.1.docx"
    build_docx(tables, docx_path)
    with pd.ExcelWriter(OUT / "PGSS_Supplementary_v0.1_DataTables.xlsx", engine="openpyxl") as xw:
        pd.DataFrame([dict(data_table=k, linked_item=v[2], rows=len(v[0]), columns=len(v[0].columns), frozen_source=v[1]) for k, v in longs.items()]) \
          .to_excel(xw, sheet_name="README", index=False)
        for k, v in longs.items():
            v[0].to_excel(xw, sheet_name=k, index=False)
    idx_rows = []
    for t in tables:
        idx_rows.append(dict(item=t["id"], group=t["group"], title=t["title"], printed_rows=len(t["df"]),
                             source_workbook="; ".join(f"{s} ({C.SRC[s].name if s in C.SRC else s})" for s, _, _ in t["src"]),
                             source_worksheet="; ".join(h for _, h, _ in t["src"]), source_rows="; ".join(str(n) for _, _, n in t["src"]),
                             source_rows_primary=t["src"][0][2] if t["src"] else np.nan,
                             source_version="frozen (" + ", ".join(sorted({s for s, _, _ in t["src"]})) + ")", fidelity=t["fidelity"],
                             data_tables="; ".join(n for n, _, _ in t["long"]), remarks=t["remarks"] or EXEMPT_NOTE.get(t["id"], "")))
    idx = pd.DataFrame(idx_rows)
    smx = source_matrix(tables)
    pend = pd.DataFrame([dict(item="ST3", content="Predictor-level missingness for the full MIMIC-IV and eICU-CRD 23-predictor pools", status="PENDING_SOURCE",
                              note="No frozen data-v5 summary table exists; R35 QC tables predate data v5 and were not used."),
                         dict(item="ST1", content="NWICU per-task distinct-patient counts", status="PENDING_SOURCE", note="Not saved in frozen task registries; cohort-level counts are in Table 1."),
                         dict(item="SM2", content="Public code/data release status and repository identifiers", status="PENDING_SOURCE", note="To be confirmed by the authors; not asserted."),
                         dict(item="ST12", content="T2 per-replicate records (internal S78)", status="SEPARATE RESOURCE", note="Synthetic records in the R47 code package; distribute with code."),
                         dict(item="—", content="Literature comparison with verbatim excerpts (internal S85)", status="INTERNAL ONLY", note="Replace with short citations; not part of the supplement.")])
    with pd.ExcelWriter(OUT / "R54_SUPP_INDEX.xlsx", engine="openpyxl") as xw:
        idx.to_excel(xw, sheet_name="1_Supplementary_index", index=False)
        smx.to_excel(xw, sheet_name="2_Source_matrix_S0_R49_R51_R52", index=False)
        pend.to_excel(xw, sheet_name="3_Pending_and_internal", index=False)
        pd.DataFrame([dict(data_table=k, linked_item=v[2], rows=len(v[0]), frozen_source=v[1]) for k, v in longs.items()]).to_excel(xw, sheet_name="4_Data_tables", index=False)
    Q = qa(tables, longs, idx, smx, docx_path)
    rec = dict(第一阶段=p1, 表数=len(tables), 长表数=len(longs), 长表总行=int(sum(len(v[0]) for v in longs.values())), QA=f"{int(Q.通过.sum())}/{len(Q)}",
               QA明细=Q.to_dict("records"), 来源矩阵=smx.original_status.value_counts().to_dict(), 决定=smx.v0_1_decision.str.split(":").str[0].value_counts().to_dict(),
               秒=round(time.time() - t0, 1))
    (OUT / "R54_QA.json").write_text(json.dumps(rec, ensure_ascii=False, indent=1, default=str), encoding="utf-8")
    pd.set_option("display.width", 220); pd.set_option("display.max_colwidth", 120)
    print(Q.to_string()); print(json.dumps({k: v for k, v in rec.items() if k not in ("第一阶段", "QA明细")}, ensure_ascii=False, indent=1))
    print(json.dumps({k: v for k, v in p1.items() if k != "改动"}, ensure_ascii=False, indent=1, default=str))
    return rec


EXEMPT_NOTE = {"ST1.1": "Full (48 task rows); merged with R49 C2 (24 ICU rows) and R51 T3 (48 rows)", "ST2.1": "Definitions table; text from R49 C3 (rows 1–6) and R48 S10",
               "ST2.2": "Pool definition read from the frozen feature-pool list used by the run code (R35_DATA_QC 4_最终特征池)",
               "ST3.1": "Condensed English rendering of R51 T4 (9 variables)", "ST3.2": "Full; phenotype rows of R51 T1 omitted (they are in ST1.1/Table 1)",
               "ST4.1": "Full for P* main-branch selections (pivot of S22 main P* + S27 + S33); all branches in data tables",
               "ST4.2": "Full; R50 O5 P* (2 rows) and HTFS (4 rows) depth tables", "ST5.2": "Full: S18 (48) + S26 (24)", "ST6.1": "Contrast definitions derived from S15 reference/candidate columns",
               "ST8.1": "Full (12 judgments)", "ST9.2": "Full (17 rows)", "ST11.1": "Full: S60 (8) merged with S61 actual shares", "ST14.3": "Full: S28 (24) + S34 (15)",
               "ST15.2": "Full (12 judgments)", "SM1.1": "Full (6 steps)"}


if __name__ == "__main__":
    main()
